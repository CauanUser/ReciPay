"""Planilha de retorno: o gestor envia os pagamentos e o supervisor atualiza os recibos (seção 4.5).

Mesmo formato da planilha de recibos (abas legíveis + aba oculta `_recipay`).
Cada supervisor recebe só os pagamentos dos próprios recibos. Importar de novo não
conta nenhum pagamento duas vezes (RN-08): a chave é o ID de cada transferência.
"""

import json
import sqlite3
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagemExcel
from openpyxl.styles import Alignment, Font

from recipay import comprovantes, troca
from recipay.auth import Usuario
from recipay.db import agora, novo_id, registrar_historico
from recipay.pagamentos import _atualizar_situacao
from recipay.pix import TIPOS_CHAVE, formatar_chave
from recipay.troca import (
    _cabecalho, _escrever_dados_ocultos, _guardar_comprovante, _ler_comprovante,
    _ler_dados_ocultos, _miniatura, _ref, _salvar_seguro, _sha256,
)

MOEDA = '"R$" #,##0.00'


@dataclass
class SupervisorRetorno:
    id: str
    nome: str
    pendentes: int        # pagamentos ainda não enviados a este supervisor
    total_centavos: int


def supervisores_para_retorno(con: sqlite3.Connection) -> list[SupervisorRetorno]:
    """Supervisores que têm pagamentos (válidos) no gestor, com quantos faltam enviar."""
    linhas = con.execute("""
        SELECT u.id, u.nome,
               SUM(CASE WHEN e.transferencia_id IS NULL THEN 1 ELSE 0 END) AS pendentes,
               SUM(CASE WHEN e.transferencia_id IS NULL THEN t.valor_centavos ELSE 0 END) AS total
        FROM transferencia t
        JOIN colaborador c ON c.id = t.colaborador_id
        JOIN usuario u ON u.id = c.supervisor_id
        LEFT JOIN retorno_envio e ON e.transferencia_id = t.id AND e.supervisor_id = u.id
        WHERE t.estornado = 0
        GROUP BY u.id ORDER BY u.nome COLLATE NOCASE
    """).fetchall()
    return [SupervisorRetorno(l["id"], l["nome"], l["pendentes"], l["total"]) for l in linhas]


def _transferencias(con, supervisor_id: str, incluir_enviados: bool) -> list[sqlite3.Row]:
    sql = """
        SELECT t.*, c.nome AS colaborador_nome, c.equipe
        FROM transferencia t JOIN colaborador c ON c.id = t.colaborador_id
        WHERE t.estornado = 0 AND c.supervisor_id = ?
    """
    if not incluir_enviados:
        sql += (" AND NOT EXISTS (SELECT 1 FROM retorno_envio e"
                " WHERE e.transferencia_id = t.id AND e.supervisor_id = c.supervisor_id)")
    sql += " ORDER BY t.data_pagamento, c.nome"
    return con.execute(sql, (supervisor_id,)).fetchall()


def exportar_retorno(
    con: sqlite3.Connection,
    usuario: Usuario,
    supervisor_id: str,
    caminho: Path | str,
    incluir_enviados: bool = False,
) -> troca.ResultadoExportacao:
    if usuario.perfil != "GESTOR":
        raise PermissionError("Somente o gestor exporta o retorno.")
    sup = con.execute("SELECT id, nome FROM usuario WHERE id = ? AND perfil = 'SUPERVISOR'",
                      (supervisor_id,)).fetchone()
    if not sup:
        raise ValueError("Supervisor não encontrado.")
    transfs = _transferencias(con, supervisor_id, incluir_enviados)
    if not transfs:
        raise ValueError(f"Não há pagamentos novos para enviar a {sup['nome']}.")

    itens, arquivos, linhas_visiveis = [], {}, []
    for t in transfs:
        conteudo = _ler_comprovante(t["comprovante_arquivo"], t["comprovante_hash"], t["colaborador_nome"])
        arquivos[t["comprovante_hash"]] = conteudo
        aplic = con.execute("""
            SELECT a.id, a.recibo_id, a.valor_centavos, r.descricao, r.data_despesa,
                   r.valor_centavos AS valor_recibo, s.saldo_centavos
            FROM aplicacao a JOIN recibo r ON r.id = a.recibo_id
            JOIN recibo_saldo s ON s.recibo_id = r.id
            WHERE a.transferencia_id = ? ORDER BY r.data_despesa""", (t["id"],)).fetchall()
        itens.append({
            "id": t["id"], "ref": _ref(t["id"]), "colaborador_id": t["colaborador_id"],
            "valor_centavos": t["valor_centavos"], "data_pagamento": t["data_pagamento"],
            "chave_pix_usada": t["chave_pix_usada"], "tipo_chave_usada": t["tipo_chave_usada"],
            "titular_usado": t["titular_usado"],
            "comprovante": {"hash": t["comprovante_hash"],
                            "extensao": Path(t["comprovante_arquivo"]).suffix.lower()},
            "aplicacoes": [{"id": a["id"], "recibo_id": a["recibo_id"],
                            "valor_centavos": a["valor_centavos"]} for a in aplic],
        })
        for a in aplic:
            linhas_visiveis.append((t, a))

    lote_id = novo_id()
    total = sum(t["valor_centavos"] for t in transfs)
    dados = {
        "versao_formato": troca.VERSAO_FORMATO, "tipo": "RETORNO", "lote_id": lote_id,
        "gerado_por": {"id": usuario.id, "nome": usuario.nome},
        "destinatario": {"id": sup["id"], "nome": sup["nome"]},
        "gerado_em": agora(), "total_centavos": total, "transferencias": itens,
    }
    texto = json.dumps(dados, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    hash_dados = _sha256(texto.encode("utf-8"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Pagamentos"
    _cabecalho(ws, ["Ref Pix", "Data do Pix", "Colaborador", "Recibo", "Abatido (R$)",
                    "Valor do Pix (R$)", "Saldo do recibo (R$)", "Chave usada"],
               [12, 12, 28, 40, 14, 16, 18, 30])
    for t, a in linhas_visiveis:
        ws.append([_ref(t["id"]), date.fromisoformat(t["data_pagamento"]), t["colaborador_nome"],
                   f"{date.fromisoformat(a['data_despesa']).strftime('%d/%m/%Y')} · {a['descricao']}",
                   a["valor_centavos"] / 100, t["valor_centavos"] / 100, a["saldo_centavos"] / 100,
                   f"{TIPOS_CHAVE.get(t['tipo_chave_usada'], '')} "
                   f"{formatar_chave(t['tipo_chave_usada'], t['chave_pix_usada'])}"])
    for linha in ws.iter_rows(min_row=2):
        linha[1].number_format = "DD/MM/YYYY"
        for celula in (linha[4], linha[5], linha[6]):
            celula.number_format = MOEDA
    ws.append([])
    ws.append(["", "", "", "TOTAL PAGO", "", total / 100])
    ws[ws.max_row][3].font = ws[ws.max_row][5].font = Font(bold=True)
    ws[ws.max_row][5].number_format = MOEDA
    ws.append([])
    ws.append([f"Retorno para {sup['nome']}, enviado por {usuario.nome} em "
               f"{date.today().strftime('%d/%m/%Y')}. Importe no Recipay; não altere os valores."])

    wf = wb.create_sheet("Comprovantes")
    _cabecalho(wf, ["Ref Pix", "Colaborador", "Valor do Pix (R$)", "Comprovante bancário"], [12, 28, 16, 40])
    for i, t in enumerate(transfs, start=2):
        wf.cell(row=i, column=1, value=_ref(t["id"]))
        wf.cell(row=i, column=2, value=t["colaborador_nome"])
        wf.cell(row=i, column=3, value=t["valor_centavos"] / 100).number_format = MOEDA
        if comprovantes.e_pdf(t["comprovante_arquivo"]):
            wf.cell(row=i, column=4, value="PDF anexado (abra no Recipay)")
        else:
            mini = _miniatura(comprovantes.caminho_absoluto(t["comprovante_arquivo"]))
            if mini:
                img = ImagemExcel(mini)
                wf.add_image(img, f"D{i}")
                wf.row_dimensions[i].height = img.height * 0.75 + 4
        for celula in wf[i]:
            celula.alignment = Alignment(vertical="top")

    _escrever_dados_ocultos(wb, texto, hash_dados, arquivos)
    caminho = Path(caminho)
    _salvar_seguro(wb, caminho)

    with con:
        con.execute(
            "INSERT INTO lote (id, tipo, gerado_por, gerado_em, hash_conteudo, qtd_itens, total_centavos)"
            " VALUES (?, 'RETORNO', ?, ?, ?, ?, ?)",
            (lote_id, usuario.id, dados["gerado_em"], hash_dados, len(transfs), total))
        for t in transfs:
            con.execute("INSERT OR REPLACE INTO retorno_envio (transferencia_id, supervisor_id, lote_id)"
                        " VALUES (?, ?, ?)", (t["id"], supervisor_id, lote_id))
            registrar_historico(con, "transferencia", t["id"], "RETORNO_EXPORTADO",
                                {"lote_id": lote_id, "supervisor": sup["nome"], "arquivo": caminho.name},
                                usuario.id, usuario.perfil)
    return troca.ResultadoExportacao(lote_id, len(transfs), total, caminho.stat().st_size)


# ================================================================ supervisor
@dataclass
class ItemRetorno:
    dados: dict
    colaborador_nome: str
    novo: bool
    recibos_desconhecidos: list[str] = field(default_factory=list)


@dataclass
class PreviaRetorno:
    arquivo_nome: str
    hash_arquivo: str
    dados: dict
    comprovantes: dict[str, bytes]
    itens: list[ItemRetorno]
    avisos: list[str] = field(default_factory=list)

    @property
    def gestor_nome(self) -> str:
        return self.dados["gerado_por"]["nome"]

    @property
    def novos(self) -> list[ItemRetorno]:
        return [i for i in self.itens if i.novo]

    @property
    def total_novos_centavos(self) -> int:
        return sum(i.dados["valor_centavos"] for i in self.novos)


def ler_retorno(con: sqlite3.Connection, usuario: Usuario, caminho: Path | str) -> PreviaRetorno:
    """Lê a planilha de retorno e monta a prévia. Não grava nada."""
    caminho = Path(caminho)
    dados, _, comps, visiveis = _ler_dados_ocultos(caminho, "Pagamentos", 5)
    if dados.get("tipo") != "RETORNO":
        raise ValueError("Este arquivo não é uma planilha de retorno de pagamentos.")
    destino = dados["destinatario"]
    if destino["id"] != usuario.id:
        raise ValueError(f"Este retorno é para {destino['nome']}, não para você.")

    for t in dados["transferencias"]:
        if visiveis.get(t["ref"]) != {t["valor_centavos"]}:
            raise ValueError(f"A planilha foi alterada no Excel (Pix {t['ref']}). "
                             "Peça um novo arquivo ao gestor.")
        h = t["comprovante"]["hash"]
        if h not in comps or _sha256(comps[h]) != h:
            raise ValueError(f"Comprovante do Pix {t['ref']} ausente ou danificado no arquivo.")

    itens = []
    for t in dados["transferencias"]:
        colab = con.execute("SELECT nome FROM colaborador WHERE id = ?", (t["colaborador_id"],)).fetchone()
        existe = con.execute("SELECT 1 FROM transferencia WHERE id = ?", (t["id"],)).fetchone()
        desconhecidos = [a["recibo_id"] for a in t["aplicacoes"]
                         if not con.execute("SELECT 1 FROM recibo WHERE id = ? AND supervisor_id = ?",
                                            (a["recibo_id"], usuario.id)).fetchone()]
        itens.append(ItemRetorno(t, colab["nome"] if colab else "(colaborador desconhecido)",
                                 not existe and not desconhecidos and colab is not None, desconhecidos))
    previa = PreviaRetorno(caminho.name, comprovantes.hash_arquivo(caminho), dados, comps, itens)
    problemas = [i for i in itens if i.recibos_desconhecidos]
    if problemas:
        previa.avisos.append(f"{len(problemas)} pagamento(s) citam recibos que não existem neste "
                             "computador e serão ignorados.")
    if not previa.novos:
        previa.avisos.append("Nenhum pagamento novo: tudo já foi importado antes.")
    return previa


@dataclass
class ResultadoRetorno:
    novos: int
    ignorados: int
    total_centavos: int
    recibos_pagos: int
    recibos_parciais: int


def importar_retorno(con: sqlite3.Connection, usuario: Usuario, previa: PreviaRetorno) -> ResultadoRetorno:
    if usuario.perfil != "SUPERVISOR":
        raise PermissionError("Somente o supervisor importa o retorno.")
    d = previa.dados
    gestor = d["gerado_por"]
    momento = agora()

    destinos = {}
    for item in previa.novos:
        comp = item.dados["comprovante"]
        if comp["hash"] not in destinos:
            destinos[comp["hash"]] = _guardar_comprovante(previa.comprovantes[comp["hash"]], comp["extensao"])

    situacoes: dict[str, str] = {}
    with con:
        existente = con.execute("SELECT perfil, local FROM usuario WHERE id = ?", (gestor["id"],)).fetchone()
        if existente and (existente["local"] or existente["perfil"] != "GESTOR"):
            raise ValueError("O arquivo indica um remetente inválido para este computador.")
        if existente:
            con.execute("UPDATE usuario SET nome = ? WHERE id = ?", (gestor["nome"], gestor["id"]))
        else:
            con.execute("INSERT INTO usuario (id, nome, perfil, local, criado_em) VALUES (?, ?, 'GESTOR', 0, ?)",
                        (gestor["id"], gestor["nome"], momento))
        if not con.execute("SELECT 1 FROM lote WHERE id = ?", (d["lote_id"],)).fetchone():
            con.execute(
                "INSERT INTO lote (id, tipo, gerado_por, gerado_em, hash_conteudo, qtd_itens, total_centavos)"
                " VALUES (?, 'RETORNO', ?, ?, ?, ?, ?)",
                (d["lote_id"], gestor["id"], d["gerado_em"], previa.hash_arquivo,
                 len(d["transferencias"]), d["total_centavos"]))

        for item in previa.novos:
            t = item.dados
            con.execute(
                "INSERT INTO transferencia (id, colaborador_id, valor_centavos, data_pagamento,"
                " chave_pix_usada, tipo_chave_usada, titular_usado, comprovante_arquivo, comprovante_hash,"
                " registrado_por, registrado_em) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (t["id"], t["colaborador_id"], t["valor_centavos"], t["data_pagamento"],
                 t["chave_pix_usada"], t["tipo_chave_usada"], t["titular_usado"],
                 destinos[t["comprovante"]["hash"]], t["comprovante"]["hash"], gestor["id"], momento))
            for a in t["aplicacoes"]:
                con.execute("INSERT INTO aplicacao (id, transferencia_id, recibo_id, valor_centavos)"
                            " VALUES (?, ?, ?, ?)", (a["id"], t["id"], a["recibo_id"], a["valor_centavos"]))
                situacoes[a["recibo_id"]] = _atualizar_situacao(con, a["recibo_id"])
                registrar_historico(con, "recibo", a["recibo_id"], "PAGAMENTO",
                                    {"transferencia_id": t["id"], "valor_centavos": a["valor_centavos"],
                                     "situacao": situacoes[a["recibo_id"]], "origem": previa.arquivo_nome},
                                    usuario.id, usuario.perfil)
        con.execute(
            "INSERT INTO importacao (id, lote_id, arquivo_nome, hash_arquivo, importado_por,"
            " importado_em, novos, ignorados) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (novo_id(), d["lote_id"], previa.arquivo_nome, previa.hash_arquivo, usuario.id, momento,
             len(previa.novos), len(previa.itens) - len(previa.novos)))

    return ResultadoRetorno(
        len(previa.novos), len(previa.itens) - len(previa.novos), previa.total_novos_centavos,
        sum(1 for s in situacoes.values() if s == "PAGO"),
        sum(1 for s in situacoes.values() if s == "PARCIAL"))


def nome_sugerido(supervisor_nome: str) -> str:
    base = "".join(ch if ch.isalnum() else "_" for ch in supervisor_nome.split()[0].lower())
    return f"retorno_{base}_{date.today().isoformat()}.xlsx"
