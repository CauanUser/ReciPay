"""Relatórios: totais do mês (por categoria, equipe e colaborador) e séries para os gráficos.

Conta os recibos pela data da despesa e ignora os cancelados.
- Lançado: valor dos recibos.
- Pago: soma dos Pix registrados (não estornados) para esses recibos.
- Pendente: o que falta pagar (saldo).
"""

import sqlite3
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from recipay.auth import Usuario
from recipay.recibos import CATEGORIAS, STATUS
from recipay.troca import _cabecalho, _salvar_seguro

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
         "setembro", "outubro", "novembro", "dezembro"]
MOEDA = '"R$" #,##0.00'
CABECALHO_FILL = PatternFill("solid", fgColor="1E5AA8")


def nome_mes(ano: int, mes: int) -> str:
    return f"{MESES[mes - 1]}/{ano}"


@dataclass
class Totais:
    lancado: int = 0
    pago: int = 0
    pendente: int = 0
    quantidade: int = 0

    def somar(self, valor: int, saldo: int) -> None:
        self.lancado += valor
        self.pago += valor - saldo
        self.pendente += saldo
        self.quantidade += 1


@dataclass
class RelatorioMes:
    ano: int
    mes: int
    geral: Totais = field(default_factory=Totais)
    por_categoria: dict[str, Totais] = field(default_factory=dict)
    por_equipe: dict[str, Totais] = field(default_factory=dict)
    por_colaborador: dict[tuple[str, str, str], Totais] = field(default_factory=dict)  # (nome, equipe, supervisor)
    linhas: list[sqlite3.Row] = field(default_factory=list)

    @property
    def titulo(self) -> str:
        return nome_mes(self.ano, self.mes)


_SQL = """
SELECT r.id, r.data_despesa, r.descricao, r.equipe, r.categoria, r.status, r.valor_centavos,
       s.saldo_centavos, c.nome AS colaborador, COALESCE(u.nome, '') AS supervisor
FROM recibo r
JOIN recibo_saldo s ON s.recibo_id = r.id
JOIN colaborador c ON c.id = r.colaborador_id
LEFT JOIN usuario u ON u.id = r.supervisor_id
WHERE r.status != 'CANCELADO' AND r.data_despesa >= ? AND r.data_despesa < ?
ORDER BY r.data_despesa, c.nome
"""


def _intervalo(ano: int, mes: int) -> tuple[str, str]:
    inicio = date(ano, mes, 1)
    fim = date(ano + (mes == 12), mes % 12 + 1, 1)
    return inicio.isoformat(), fim.isoformat()


def do_mes(con: sqlite3.Connection, ano: int, mes: int) -> RelatorioMes:
    rel = RelatorioMes(ano, mes)
    rel.linhas = con.execute(_SQL, _intervalo(ano, mes)).fetchall()
    for l in rel.linhas:
        v, s = l["valor_centavos"], max(l["saldo_centavos"], 0)
        rel.geral.somar(v, s)
        rel.por_categoria.setdefault(l["categoria"] or "OUTROS", Totais()).somar(v, s)
        rel.por_equipe.setdefault(l["equipe"], Totais()).somar(v, s)
        rel.por_colaborador.setdefault((l["colaborador"], l["equipe"], l["supervisor"]), Totais()).somar(v, s)
    return rel


def meses_com_dados(con: sqlite3.Connection) -> list[tuple[int, int]]:
    """Meses (ano, mês) com recibos, do mais recente para o mais antigo. Sempre inclui o atual."""
    hoje = date.today()
    meses = {(hoje.year, hoje.month)}
    for (am,) in con.execute("SELECT DISTINCT substr(data_despesa, 1, 7) FROM recibo WHERE status != 'CANCELADO'"):
        meses.add((int(am[:4]), int(am[5:7])))
    return sorted(meses, reverse=True)


def ultimos_meses(con: sqlite3.Connection, quantidade: int = 6) -> list[tuple[str, int]]:
    """[(rótulo 'out/26', total lançado)] dos últimos meses, do mais antigo ao atual."""
    hoje = date.today()
    resultado = []
    for i in range(quantidade - 1, -1, -1):
        total_meses = hoje.year * 12 + hoje.month - 1 - i
        ano, mes = divmod(total_meses, 12)
        mes += 1
        total = con.execute(
            "SELECT COALESCE(SUM(valor_centavos), 0) FROM recibo WHERE status != 'CANCELADO'"
            " AND data_despesa >= ? AND data_despesa < ?", _intervalo(ano, mes)).fetchone()[0]
        resultado.append((f"{MESES[mes - 1][:3]}/{str(ano)[2:]}", total))
    return resultado


def top_equipes(con: sqlite3.Connection, meses: int = 6, limite: int = 5) -> list[tuple[str, int]]:
    """Equipes que mais gastaram nos últimos meses (inclui o atual)."""
    hoje = date.today()
    total_meses = hoje.year * 12 + hoje.month - 1 - (meses - 1)
    inicio = date(total_meses // 12, total_meses % 12 + 1, 1).isoformat()
    return [(e, v) for e, v in con.execute(
        "SELECT equipe, SUM(valor_centavos) AS v FROM recibo WHERE status != 'CANCELADO'"
        " AND data_despesa >= ? GROUP BY equipe ORDER BY v DESC LIMIT ?", (inicio, limite))]


def _escrever_tabela(ws, linha0: int, titulo: str, rotulo: str, dados: list[tuple[str, Totais]]) -> int:
    ws.cell(row=linha0, column=1, value=titulo).font = Font(bold=True, size=12)
    cab = linha0 + 1
    for j, t in enumerate([rotulo, "Recibos", "Lançado (R$)", "Pago (R$)", "Pendente (R$)"], start=1):
        c = ws.cell(row=cab, column=j, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = CABECALHO_FILL
    linha = cab + 1
    for nome, t in dados:
        for j, v in enumerate([nome, t.quantidade, t.lancado / 100, t.pago / 100, t.pendente / 100], start=1):
            c = ws.cell(row=linha, column=j, value=v)
            if j >= 3:
                c.number_format = MOEDA
        linha += 1
    return linha


def exportar_excel(con: sqlite3.Connection, usuario: Usuario, ano: int, mes: int, caminho: Path | str) -> RelatorioMes:
    rel = do_mes(con, ano, mes)
    wb = Workbook()

    # ---- Resumo
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Recipay — Relatório de {rel.titulo}"
    ws["A1"].font = Font(bold=True, size=14, color="1E5AA8")
    ws["A2"] = (f"Gerado por {usuario.nome} ({usuario.perfil_nome}) em {date.today():%d/%m/%Y}. "
                "Recibos pela data da despesa; cancelados não entram.")
    for i, (rotulo, valor) in enumerate((("Recibos", rel.geral.quantidade), ("Lançado", rel.geral.lancado),
                                          ("Pago", rel.geral.pago), ("Pendente", rel.geral.pendente))):
        ws.cell(row=4 + i, column=1, value=rotulo).font = Font(bold=True)
        c = ws.cell(row=4 + i, column=3, value=valor if i == 0 else valor / 100)
        c.font = Font(size=12, bold=True)
        if i:
            c.number_format = MOEDA

    categorias = sorted(rel.por_categoria.items(), key=lambda x: -x[1].lancado)
    inicio_cat = 9
    fim_cat = _escrever_tabela(ws, inicio_cat, "Por categoria", "Categoria",
                               [(CATEGORIAS.get(k, k), t) for k, t in categorias])
    equipes = sorted(rel.por_equipe.items(), key=lambda x: -x[1].lancado)
    inicio_eq = fim_cat + 2
    fim_eq = _escrever_tabela(ws, inicio_eq, "Por equipe", "Equipe", equipes)
    for col, larg in zip("ABCDE", (34, 10, 16, 16, 16)):
        ws.column_dimensions[col].width = larg

    # Impressão: paisagem, cabendo na largura da folha
    for planilha in (ws,):
        planilha.page_setup.orientation = "landscape"
        planilha.page_setup.fitToWidth = 1
        planilha.page_setup.fitToHeight = 0
        planilha.sheet_properties.pageSetUpPr.fitToPage = True

    # Gráficos nativos do Excel (uma série cada: o valor lançado)
    for titulo, ini, fim, ancora in (("Lançado por categoria", inicio_cat, fim_cat, "G2"),
                                     ("Lançado por equipe", inicio_eq, fim_eq, "G18")):
        if fim - ini <= 2:
            continue
        grafico = BarChart()
        grafico.type = "bar"
        grafico.title = titulo
        grafico.legend = None
        grafico.x_axis.scaling.orientation = "maxMin"  # maior em cima, como na tabela
        grafico.x_axis.delete = False
        grafico.y_axis.delete = False
        grafico.height, grafico.width = 7.5, 15
        grafico.add_data(Reference(ws, min_col=3, min_row=ini + 1, max_row=fim - 1), titles_from_data=True)
        grafico.set_categories(Reference(ws, min_col=1, min_row=ini + 2, max_row=fim - 1))
        grafico.series[0].graphicalProperties.solidFill = "1E5AA8"
        grafico.series[0].graphicalProperties.line.noFill = True
        ws.add_chart(grafico, ancora)

    # ---- Por colaborador
    wc = wb.create_sheet("Por colaborador")
    _cabecalho(wc, ["Colaborador", "Equipe", "Supervisor", "Recibos", "Lançado (R$)", "Pago (R$)", "Pendente (R$)"],
               [32, 18, 24, 10, 16, 16, 16])
    for (nome, equipe, sup), t in sorted(rel.por_colaborador.items(), key=lambda x: -x[1].lancado):
        wc.append([nome, equipe, sup, t.quantidade, t.lancado / 100, t.pago / 100, t.pendente / 100])
    for linha in wc.iter_rows(min_row=2):
        for c in linha[4:7]:
            c.number_format = MOEDA
    wc.auto_filter.ref = wc.dimensions

    # ---- Recibos (detalhe)
    wr = wb.create_sheet("Recibos")
    _cabecalho(wr, ["Data", "Colaborador", "Equipe", "Supervisor", "Categoria", "Descrição", "Valor (R$)",
                    "Pago (R$)", "Pendente (R$)", "Situação"], [12, 28, 16, 22, 26, 36, 14, 14, 14, 12])
    for l in rel.linhas:
        saldo = max(l["saldo_centavos"], 0)
        wr.append([date.fromisoformat(l["data_despesa"]), l["colaborador"], l["equipe"], l["supervisor"],
                   CATEGORIAS.get(l["categoria"], "Outros"), l["descricao"], l["valor_centavos"] / 100,
                   (l["valor_centavos"] - saldo) / 100, saldo / 100, STATUS.get(l["status"], l["status"])])
    for linha in wr.iter_rows(min_row=2):
        linha[0].number_format = "DD/MM/YYYY"
        for c in linha[6:9]:
            c.number_format = MOEDA
    wr.auto_filter.ref = wr.dimensions

    _salvar_seguro(wb, Path(caminho))
    return rel
