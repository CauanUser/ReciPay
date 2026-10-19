"""Importação de colaboradores por planilha .xlsx (com prévia antes de gravar)."""

import sqlite3
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from recipay import colaboradores
from recipay.auth import Usuario
from recipay.pix import TIPOS_CHAVE, detectar_tipo, formatar_chave, normalizar_tipo

COLUNAS = ["nome", "equipe", "chave_pix", "tipo_chave", "titular_pix"]
OBRIGATORIAS = ["nome", "equipe", "chave_pix"]
MAX_LINHAS = 5000


@dataclass
class LinhaPrevia:
    numero: int                      # número da linha na planilha
    original: dict
    dados: dict | None = None        # dados normalizados, se válidos
    erros: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    @property
    def valida(self) -> bool:
        return not self.erros


def _cabecalho(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    t = t.strip().lower().replace(" ", "_").replace("-", "_")
    apelidos = {"chave": "chave_pix", "pix": "chave_pix", "tipo": "tipo_chave",
                "titular": "titular_pix", "tipo_da_chave": "tipo_chave",
                "nome_completo": "nome", "time": "equipe"}
    return apelidos.get(t, t)


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # números grandes do Excel (CPF, telefone) chegam como float
    return str(valor).strip()


def ler(con: sqlite3.Connection, caminho: Path | str) -> list[LinhaPrevia]:
    """Lê a planilha e devolve a prévia, sem gravar nada."""
    try:
        wb = load_workbook(caminho, read_only=True, data_only=True)
    except Exception as e:  # arquivo corrompido, protegido por senha etc.
        raise ValueError(f"Não foi possível abrir a planilha: {e}") from e
    try:
        return _ler_previa(con, wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()  # libera o arquivo (no Windows, evita deixá-lo travado)


def _ler_previa(con: sqlite3.Connection, linhas) -> list[LinhaPrevia]:
    cab = next(linhas, None)
    if not cab:
        raise ValueError("A planilha está vazia.")
    indices = {}
    for i, titulo in enumerate(cab):
        nome = _cabecalho(titulo)
        if nome in COLUNAS and nome not in indices:
            indices[nome] = i
    faltando = [c for c in OBRIGATORIAS if c not in indices]
    if faltando:
        raise ValueError(
            "Colunas obrigatórias não encontradas na primeira linha: " + ", ".join(faltando)
            + ". Use o modelo de planilha."
        )

    previa: list[LinhaPrevia] = []
    chaves_na_planilha: dict[str, int] = {}
    for numero, valores in enumerate(linhas, start=2):
        original = {c: _texto(valores[i]) if i < len(valores) else "" for c, i in indices.items()}
        if not any(original.values()):
            continue  # linha em branco
        if len(previa) >= MAX_LINHAS:
            raise ValueError(f"A planilha tem mais de {MAX_LINHAS} linhas.")
        linha = LinhaPrevia(numero, original)
        previa.append(linha)

        tipo_txt = original.get("tipo_chave", "")
        if tipo_txt:
            tipo = normalizar_tipo(tipo_txt)
            if not tipo:
                linha.erros.append(f"Tipo de chave desconhecido: {tipo_txt}.")
                continue
        else:
            tipo = detectar_tipo(original["chave_pix"])
            if not tipo:
                linha.erros.append("Informe o tipo da chave (não foi possível identificar).")
                continue
            linha.avisos.append(f"Tipo identificado: {TIPOS_CHAVE[tipo]}.")

        titular = original.get("titular_pix", "")
        if not titular:
            titular = original["nome"]
            linha.avisos.append("Titular vazio: usado o nome do colaborador.")

        try:
            dados = colaboradores.validar({**original, "tipo_chave": tipo, "titular_pix": titular})
        except ValueError as e:
            linha.erros.append(str(e))
            continue

        if dados["chave_pix"] in chaves_na_planilha:
            linha.erros.append(f"Chave repetida na linha {chaves_na_planilha[dados['chave_pix']]}.")
            continue
        chaves_na_planilha[dados["chave_pix"]] = numero

        existente = colaboradores.chave_em_uso(con, dados["chave_pix"])
        if existente:
            linha.erros.append(f"Chave já cadastrada para {existente.nome}.")
            continue
        linha.dados = dados

    if not previa:
        raise ValueError("Nenhuma linha com dados encontrada.")
    return previa


def gravar(con: sqlite3.Connection, usuario: Usuario, previa: list[LinhaPrevia]) -> int:
    """Grava somente as linhas válidas. Devolve quantas foram importadas."""
    total = 0
    for linha in previa:
        if linha.valida and linha.dados:
            colaboradores.criar(con, usuario, linha.dados, origem="PLANILHA")
            total += 1
    return total


def _estilizar_cabecalho(ws, larguras: tuple[int, ...]) -> None:
    azul = PatternFill("solid", fgColor="1E5AA8")
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = azul
    for indice, largura in enumerate(larguras):
        ws.column_dimensions[chr(ord("A") + indice)].width = largura
    ws.freeze_panes = "A2"


def exportar(caminho: Path | str, lista: list[colaboradores.Colaborador]) -> None:
    """Exporta colaboradores no mesmo formato da importação (pode ser reimportada)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Colaboradores"
    ws.append(COLUNAS + ["situacao"])
    for c in lista:
        ws.append([c.nome, c.equipe, formatar_chave(c.tipo_chave, c.chave_pix),
                   TIPOS_CHAVE[c.tipo_chave], c.titular_pix, "Ativo" if c.ativo else "Inativo"])
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=3).number_format = "@"
    _estilizar_cabecalho(ws, (32, 18, 38, 16, 32, 12))
    ws.auto_filter.ref = ws.dimensions
    wb.save(caminho)


def salvar_modelo(caminho: Path | str) -> None:
    """Gera uma planilha modelo com cabeçalho, exemplo e instruções."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Colaboradores"
    ws.append(COLUNAS)
    ws.append(["Maria da Silva", "Equipe Norte", "maria@exemplo.com", "E-mail", "Maria da Silva"])
    azul = PatternFill("solid", fgColor="1E5AA8")
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = azul
    for coluna, largura in zip("ABCDE", (32, 18, 38, 16, 32)):
        ws.column_dimensions[coluna].width = largura
    # Coluna da chave como texto, para o Excel não cortar zeros do CPF.
    for linha in range(2, 1001):
        ws.cell(row=linha, column=3).number_format = "@"

    ajuda = wb.create_sheet("Instruções")
    for texto in (
        "Preencha uma linha por colaborador na aba Colaboradores. Apague a linha de exemplo.",
        "Obrigatórias: nome, equipe, chave_pix.",
        "tipo_chave: CPF, CNPJ, E-mail, Telefone ou Aleatória. Se vazio, o Recipay tenta identificar.",
        "titular_pix: nome do titular da conta. Se vazio, usa o nome do colaborador.",
        "Telefone: DDD + número, ex.: (11) 99999-8888.",
        "Antes de gravar, o Recipay mostra uma prévia com os erros de cada linha.",
    ):
        ajuda.append([texto])
    ajuda.column_dimensions["A"].width = 100
    wb.save(caminho)
