"""Troca de arquivos entre supervisor e gestor por planilha Excel (D-07, seção 8).

A planilha tem abas legíveis (Recibos, Colaboradores, Comprovantes) e uma aba
oculta `_recipay` com os dados completos em JSON e os comprovantes em base64.
A importação usa só a aba oculta e confere se as abas visíveis não foram alteradas.
"""

import base64
import hashlib
import io
import json
import os
import sqlite3
import tempfile
import uuid
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ImagemExcel
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

from recipay import comprovantes, config, recibos
from recipay.auth import Usuario
from recipay.db import agora, novo_id, registrar_historico

VERSAO_FORMATO = 1
ABA_DADOS = "_recipay"
MARCA = "RECIPAY"
TAMANHO_PEDACO = 32000          # limite do Excel: 32.767 caracteres por célula
LIMITE_AVISO_BYTES = 20 * 1024 * 1024
AZUL = "1E5AA8"


# =============================================================== utilitários
def _sha256(dados: bytes) -> str:
    return hashlib.sha256(dados).hexdigest()


def _ref(recibo_id: str) -> str:
    """Referência curta para cruzar as abas visíveis (ex.: 3F2A9C1B)."""
    return recibo_id.replace("-", "")[:8].upper()


def _cabecalho(ws, titulos: list[str], larguras: list[int]) -> None:
    ws.append(titulos)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=AZUL)
        celula.alignment = Alignment(vertical="center")
    for i, largura in enumerate(larguras):
        ws.column_dimensions[chr(ord("A") + i)].width = largura
    ws.freeze_panes = "A2"


def _miniatura(caminho: Path) -> io.BytesIO | None:
    try:
        with Image.open(caminho) as img:
            img = img.convert("RGB")
            img.thumbnail((260, 260))
            saida = io.BytesIO()
            img.save(saida, "JPEG", quality=80)
            saida.seek(0)
            return saida
    except OSError:
        return None


# =============================================================== exportação
@dataclass
class ResultadoExportacao:
    lote_id: str
    quantidade: int
    total_centavos: int
    tamanho_bytes: int

    @property
    def grande_demais(self) -> bool:
        return self.tamanho_bytes > LIMITE_AVISO_BYTES


def exportar_recibos(
    con: sqlite3.Connection, usuario: Usuario, recibo_ids: list[str], caminho: Path | str
) -> ResultadoExportacao:
    """Gera a planilha para o gestor e marca os recibos como ENVIADO.

    Aceita rascunhos e também recibos já enviados (reenvio, caso o arquivo se perca):
    os IDs são os mesmos, então o gestor não importa nada em dobro.
    """
    if usuario.perfil != "SUPERVISOR":
        raise PermissionError("Somente o supervisor exporta recibos.")
    if not recibo_ids:
        raise ValueError("Selecione pelo menos um recibo.")

    lista = []
    for rid in recibo_ids:
        r = recibos.buscar(con, rid)
        if not r or r.supervisor_id != usuario.id:
            raise ValueError("Recibo não encontrado.")
        if r.status not in ("RASCUNHO", "ENVIADO"):
            raise ValueError(f"O recibo de {r.colaborador_nome} não pode ser exportado ({r.status}).")
        lista.append(r)
    lista.sort(key=lambda r: (r.colaborador_nome.lower(), r.data_despesa))

    # Colaboradores envolvidos (o gestor precisa da chave Pix para pagar).
    colabs = {}
    for r in lista:
        if r.colaborador_id not in colabs:
            c = con.execute("SELECT * FROM colaborador WHERE id = ?", (r.colaborador_id,)).fetchone()
            colabs[c["id"]] = {k: c[k] for k in
                               ("id", "nome", "equipe", "chave_pix", "tipo_chave", "titular_pix")}

    # Comprovantes: confere a integridade de cada um antes de exportar.
    arquivos_comprov: dict[str, bytes] = {}
    itens = []
    for r in lista:
        caminho_comp = comprovantes.caminho_absoluto(r.comprovante_arquivo)
        try:
            conteudo = caminho_comp.read_bytes()
        except OSError as e:
            raise ValueError(f"Comprovante de {r.colaborador_nome} não encontrado no computador.") from e
        hash_ = _sha256(conteudo)
        esperado = con.execute("SELECT comprovante_hash FROM recibo WHERE id = ?", (r.id,)).fetchone()[0]
        if hash_ != esperado:
            raise ValueError(f"O comprovante de {r.colaborador_nome} foi alterado fora do programa.")
        arquivos_comprov[hash_] = conteudo
        itens.append({
            "id": r.id, "ref": _ref(r.id), "colaborador_id": r.colaborador_id, "equipe": r.equipe,
            "valor_centavos": r.valor_centavos, "descricao": r.descricao,
            "data_despesa": r.data_despesa,
            "comprovante": {"hash": hash_, "extensao": Path(r.comprovante_arquivo).suffix.lower()},
        })

    lote_id = novo_id()
    total = sum(r.valor_centavos for r in lista)
    dados = {
        "versao_formato": VERSAO_FORMATO,
        "tipo": "RECIBOS",
        "lote_id": lote_id,
        "gerado_por": {"id": usuario.id, "nome": usuario.nome},
        "gerado_em": agora(),
        "total_centavos": total,
        "colaboradores": list(colabs.values()),
        "recibos": itens,
    }
    texto_dados = json.dumps(dados, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    hash_dados = _sha256(texto_dados.encode("utf-8"))

    wb = _montar_planilha(lista, colabs, dados, texto_dados, hash_dados, arquivos_comprov)

    caminho = Path(caminho)
    _salvar_seguro(wb, caminho)

    with con:
        con.execute(
            "INSERT INTO lote (id, tipo, gerado_por, gerado_em, hash_conteudo, qtd_itens, total_centavos)"
            " VALUES (?, 'RECIBOS', ?, ?, ?, ?, ?)",
            (lote_id, usuario.id, dados["gerado_em"], hash_dados, len(lista), total),
        )
        for r in lista:
            con.execute("UPDATE recibo SET status = 'ENVIADO', lote_id = ? WHERE id = ?",
                        (lote_id, r.id))
            registrar_historico(con, "recibo", r.id,
                                "REENVIADO" if r.status == "ENVIADO" else "EXPORTADO",
                                {"lote_id": lote_id, "arquivo": caminho.name},
                                usuario.id, usuario.perfil)
    return ResultadoExportacao(lote_id, len(lista), total, caminho.stat().st_size)


def _escrever_dados_ocultos(wb: Workbook, texto_dados: str, hash_dados: str,
                            arquivos: dict[str, bytes]) -> None:
    """Aba oculta com os dados em JSON e os comprovantes em base64, em pedaços."""
    wd = wb.create_sheet(ABA_DADOS)
    wd.sheet_state = "veryHidden"
    wd.append([MARCA, VERSAO_FORMATO, "Não altere esta aba."])
    wd.append(["hash", 0, hash_dados])
    for chave, texto in [("dados", texto_dados)] + [
        (f"comp:{h}", base64.b64encode(b).decode("ascii")) for h, b in arquivos.items()
    ]:
        for n, inicio in enumerate(range(0, len(texto), TAMANHO_PEDACO)):
            wd.append([chave, n, texto[inicio:inicio + TAMANHO_PEDACO]])


def _salvar_seguro(wb: Workbook, caminho: Path) -> None:
    """Grava num temporário e depois troca, para não deixar arquivo pela metade."""
    fd, temporario = tempfile.mkstemp(suffix=".xlsx", dir=caminho.parent)
    os.close(fd)
    try:
        wb.save(temporario)
        os.replace(temporario, caminho)
    except OSError as e:
        Path(temporario).unlink(missing_ok=True)
        raise OSError(f"Não foi possível salvar a planilha: {e}. Ela está aberta no Excel?") from e


def _ler_comprovante(arquivo: str, hash_esperado: str, quem: str) -> bytes:
    try:
        conteudo = comprovantes.caminho_absoluto(arquivo).read_bytes()
    except OSError as e:
        raise ValueError(f"Comprovante de {quem} não encontrado no computador.") from e
    if _sha256(conteudo) != hash_esperado:
        raise ValueError(f"O comprovante de {quem} foi alterado fora do programa.")
    return conteudo


def _guardar_comprovante(conteudo: bytes, extensao: str) -> str:
    destino = config.pasta_comprovantes() / f"{uuid.uuid4().hex}{extensao}"
    destino.write_bytes(conteudo)
    return destino.relative_to(config.pasta_dados()).as_posix()


def _montar_planilha(lista, colabs, dados, texto_dados, hash_dados, arquivos_comprov) -> Workbook:
    from recipay.pix import TIPOS_CHAVE, formatar_chave

    wb = Workbook()
    # ---- Recibos (legível)
    ws = wb.active
    ws.title = "Recibos"
    _cabecalho(ws, ["Ref", "Data", "Colaborador", "Equipe", "Descrição", "Valor (R$)"],
               [12, 12, 30, 16, 40, 14])
    for r in lista:
        ws.append([_ref(r.id), date.fromisoformat(r.data_despesa), r.colaborador_nome,
                   r.equipe, r.descricao, r.valor_centavos / 100])
    for linha in ws.iter_rows(min_row=2):
        linha[1].number_format = "DD/MM/YYYY"
        linha[5].number_format = '"R$" #,##0.00'
    ws.append([])
    ws.append(["", "", "", "", "TOTAL", dados["total_centavos"] / 100])
    total = ws[ws.max_row]
    total[4].font = total[5].font = Font(bold=True)
    total[5].number_format = '"R$" #,##0.00'
    ws.append([])
    ws.append([f"Enviado por {dados['gerado_por']['nome']} em "
               f"{date.fromisoformat(dados['gerado_em'][:10]).strftime('%d/%m/%Y')}. Importe esta planilha no Recipay; não altere os valores."])

    # ---- Colaboradores (legível)
    wc = wb.create_sheet("Colaboradores")
    _cabecalho(wc, ["Nome", "Equipe", "Tipo da chave", "Chave Pix", "Titular"], [30, 16, 16, 36, 30])
    for c in colabs.values():
        wc.append([c["nome"], c["equipe"], TIPOS_CHAVE[c["tipo_chave"]],
                   formatar_chave(c["tipo_chave"], c["chave_pix"]), c["titular_pix"]])

    # ---- Comprovantes (miniaturas)
    wf = wb.create_sheet("Comprovantes")
    _cabecalho(wf, ["Ref", "Colaborador", "Valor (R$)", "Comprovante"], [12, 30, 14, 40])
    for i, r in enumerate(lista, start=2):
        wf.cell(row=i, column=1, value=_ref(r.id))
        wf.cell(row=i, column=2, value=r.colaborador_nome)
        wf.cell(row=i, column=3, value=r.valor_centavos / 100).number_format = '"R$" #,##0.00'
        if comprovantes.e_pdf(r.comprovante_arquivo):
            wf.cell(row=i, column=4, value="PDF anexado (abra no Recipay)")
        else:
            mini = _miniatura(comprovantes.caminho_absoluto(r.comprovante_arquivo))
            if mini:
                imagem = ImagemExcel(mini)
                wf.add_image(imagem, f"D{i}")
                wf.row_dimensions[i].height = imagem.height * 0.75 + 4
    for linha in wf.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(vertical="top")

    _escrever_dados_ocultos(wb, texto_dados, hash_dados, arquivos_comprov)
    return wb


# =============================================================== leitura
@dataclass
class ItemPrevia:
    dados: dict
    colaborador_nome: str
    novo: bool


@dataclass
class ChaveAlterada:
    colaborador_id: str
    nome: str
    atual: str
    tipo_atual: str
    nova: str
    tipo_nova: str


@dataclass
class PreviaImportacao:
    arquivo_nome: str
    hash_arquivo: str
    dados: dict
    comprovantes: dict[str, bytes]
    itens: list[ItemPrevia]
    colaboradores_novos: list[dict]
    chaves_alteradas: list[ChaveAlterada]
    lote_ja_importado_em: str | None = None
    avisos: list[str] = field(default_factory=list)

    @property
    def supervisor_nome(self) -> str:
        return self.dados["gerado_por"]["nome"]

    @property
    def novos(self) -> list[ItemPrevia]:
        return [i for i in self.itens if i.novo]

    @property
    def total_novos_centavos(self) -> int:
        return sum(i.dados["valor_centavos"] for i in self.novos)


def _ler_dados_ocultos(
    caminho: Path, aba_visivel: str = "Recibos", coluna_valor: int = 5
) -> tuple[dict, str, dict[str, bytes], dict[str, set[int]]]:
    try:
        wb = load_workbook(caminho, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Não foi possível abrir a planilha: {e}") from e
    try:
        if ABA_DADOS not in wb.sheetnames:
            raise ValueError("Esta planilha não foi gerada pelo Recipay (ou foi salva de novo "
                             "por outro programa e perdeu os dados). Peça o arquivo original.")
        pedacos: dict[str, list[tuple[int, str]]] = {}
        linhas = wb[ABA_DADOS].iter_rows(values_only=True)
        cab = next(linhas, None)
        if not cab or cab[0] != MARCA:
            raise ValueError("Planilha do Recipay danificada (marca não encontrada).")
        if int(cab[1] or 0) > VERSAO_FORMATO:
            raise ValueError("Esta planilha foi gerada por uma versão mais nova do Recipay. "
                             "Atualize o programa.")
        for linha in linhas:
            if not linha or len(linha) < 3 or linha[0] is None:
                continue
            pedacos.setdefault(str(linha[0]), []).append((int(linha[1]), str(linha[2] or "")))
        juntar = {k: "".join(t for _, t in sorted(v)) for k, v in pedacos.items()}

        # Valores da aba visível, para detectar edição no Excel.
        visiveis: dict[str, set[int]] = {}
        if aba_visivel in wb.sheetnames:
            for linha in wb[aba_visivel].iter_rows(min_row=2, values_only=True):
                if (linha and len(linha) > coluna_valor and linha[0]
                        and isinstance(linha[coluna_valor], (int, float))):
                    visiveis.setdefault(str(linha[0]), set()).add(round(linha[coluna_valor] * 100))
    finally:
        wb.close()

    texto = juntar.get("dados")
    if not texto:
        raise ValueError("Planilha do Recipay danificada (dados não encontrados).")
    if _sha256(texto.encode("utf-8")) != juntar.get("hash"):
        raise ValueError("A planilha foi alterada depois de exportada. Peça um novo arquivo ao supervisor.")
    dados = json.loads(texto)
    comps = {}
    for chave, b64 in juntar.items():
        if chave.startswith("comp:"):
            comps[chave[5:]] = base64.b64decode(b64)
    return dados, texto, comps, visiveis


def ler_arquivo(con: sqlite3.Connection, caminho: Path | str) -> PreviaImportacao:
    """Lê a planilha do supervisor e monta a prévia. Não grava nada."""
    caminho = Path(caminho)
    dados, _, comps, visiveis = _ler_dados_ocultos(caminho)
    if dados.get("tipo") != "RECIBOS":
        raise ValueError("Este arquivo não é uma planilha de recibos.")

    # Abas visíveis alteradas?
    for item in dados["recibos"]:
        if visiveis.get(item["ref"]) != {item["valor_centavos"]}:
            raise ValueError(
                f"A planilha foi alterada no Excel (recibo {item['ref']}). "
                "Os valores visíveis não batem com os registrados. Peça um novo arquivo ao supervisor.")

    # Comprovantes íntegros?
    for item in dados["recibos"]:
        h = item["comprovante"]["hash"]
        if h not in comps or _sha256(comps[h]) != h:
            raise ValueError(f"Comprovante do recibo {item['ref']} ausente ou danificado no arquivo.")

    colabs = {c["id"]: c for c in dados["colaboradores"]}
    itens = []
    for item in dados["recibos"]:
        existe = con.execute("SELECT 1 FROM recibo WHERE id = ?", (item["id"],)).fetchone()
        itens.append(ItemPrevia(item, colabs[item["colaborador_id"]]["nome"], not existe))

    novos_colabs, alteradas = [], []
    for c in colabs.values():
        atual = con.execute("SELECT * FROM colaborador WHERE id = ?", (c["id"],)).fetchone()
        if not atual:
            novos_colabs.append(c)
        elif (atual["chave_pix"], atual["tipo_chave"]) != (c["chave_pix"], c["tipo_chave"]):
            alteradas.append(ChaveAlterada(c["id"], c["nome"], atual["chave_pix"],
                                           atual["tipo_chave"], c["chave_pix"], c["tipo_chave"]))

    lote = con.execute(
        "SELECT importado_em FROM importacao WHERE lote_id = ? ORDER BY importado_em LIMIT 1",
        (dados["lote_id"],)).fetchone()
    hash_arquivo = comprovantes.hash_arquivo(caminho)
    previa = PreviaImportacao(caminho.name, hash_arquivo, dados, comps, itens, novos_colabs,
                              alteradas, lote["importado_em"] if lote else None)
    if not previa.novos:
        previa.avisos.append("Todos os recibos desta planilha já foram importados antes.")
    return previa


# =============================================================== importação
@dataclass
class ResultadoImportacao:
    novos: int
    ignorados: int
    total_centavos: int


def importar(
    con: sqlite3.Connection,
    usuario: Usuario,
    previa: PreviaImportacao,
    usar_chave_nova: dict[str, bool] | None = None,
) -> ResultadoImportacao:
    """Grava os recibos novos como PENDENTE. Os já importados são ignorados (RN-08).

    usar_chave_nova: para cada colaborador com chave alterada, True = usar a da planilha.
    """
    if usuario.perfil != "GESTOR":
        raise PermissionError("Somente o gestor importa recibos.")
    usar_chave_nova = usar_chave_nova or {}
    d = previa.dados
    sup = d["gerado_por"]
    momento = agora()

    # Comprovantes primeiro (arquivos). Só os dos recibos novos.
    destinos = {}
    for item in previa.novos:
        comp = item.dados["comprovante"]
        if comp["hash"] in destinos:
            continue
        nome = f"{uuid.uuid4().hex}{comp['extensao']}"
        destino = config.pasta_comprovantes() / nome
        destino.write_bytes(previa.comprovantes[comp["hash"]])
        destinos[comp["hash"]] = destino.relative_to(config.pasta_dados()).as_posix()

    with con:
        # Supervisor de origem (sem senha: não faz login neste PC).
        existente = con.execute("SELECT perfil, local FROM usuario WHERE id = ?", (sup["id"],)).fetchone()
        if existente and (existente["local"] or existente["perfil"] != "SUPERVISOR"):
            raise ValueError("O arquivo indica um remetente inválido para este computador.")
        if existente:
            con.execute("UPDATE usuario SET nome = ? WHERE id = ?", (sup["nome"], sup["id"]))
        else:
            con.execute("INSERT INTO usuario (id, nome, perfil, local, criado_em)"
                        " VALUES (?, ?, 'SUPERVISOR', 0, ?)", (sup["id"], sup["nome"], momento))

        # Colaboradores
        for c in d["colaboradores"]:
            atual = con.execute("SELECT id FROM colaborador WHERE id = ?", (c["id"],)).fetchone()
            if not atual:
                con.execute(
                    "INSERT INTO colaborador (id, supervisor_id, nome, equipe, chave_pix, tipo_chave,"
                    " titular_pix, ativo, criado_em, atualizado_em) VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?, ?)",
                    (c["id"], sup["id"], c["nome"], c["equipe"], c["chave_pix"], c["tipo_chave"],
                     c["titular_pix"], momento, momento))
                registrar_historico(con, "colaborador", c["id"], "IMPORTADO",
                                    {"lote_id": d["lote_id"]}, usuario.id, usuario.perfil)
            else:
                usar_nova = usar_chave_nova.get(c["id"], True)
                if usar_nova:
                    con.execute(
                        "UPDATE colaborador SET nome = ?, equipe = ?, chave_pix = ?, tipo_chave = ?,"
                        " titular_pix = ?, atualizado_em = ? WHERE id = ?",
                        (c["nome"], c["equipe"], c["chave_pix"], c["tipo_chave"], c["titular_pix"],
                         momento, c["id"]))
                alterada = next((a for a in previa.chaves_alteradas if a.colaborador_id == c["id"]), None)
                if alterada:
                    registrar_historico(
                        con, "colaborador", c["id"],
                        "CHAVE_ATUALIZADA" if usar_nova else "CHAVE_MANTIDA",
                        {"antes": alterada.atual, "planilha": alterada.nova}, usuario.id, usuario.perfil)

        # Lote
        if not con.execute("SELECT 1 FROM lote WHERE id = ?", (d["lote_id"],)).fetchone():
            con.execute(
                "INSERT INTO lote (id, tipo, gerado_por, gerado_em, hash_conteudo, qtd_itens, total_centavos)"
                " VALUES (?, 'RECIBOS', ?, ?, ?, ?, ?)",
                (d["lote_id"], sup["id"], d["gerado_em"], previa.hash_arquivo,
                 len(d["recibos"]), d["total_centavos"]))

        # Recibos novos
        for item in previa.novos:
            r = item.dados
            con.execute(
                "INSERT INTO recibo (id, colaborador_id, supervisor_id, equipe, valor_centavos,"
                " descricao, data_despesa, comprovante_arquivo, comprovante_hash, status, lote_id, criado_em)"
                " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'PENDENTE', ?, ?)",
                (r["id"], r["colaborador_id"], sup["id"], r["equipe"], r["valor_centavos"],
                 r["descricao"], r["data_despesa"], destinos[r["comprovante"]["hash"]],
                 r["comprovante"]["hash"], d["lote_id"], momento))
            registrar_historico(con, "recibo", r["id"], "IMPORTADO",
                                {"lote_id": d["lote_id"], "arquivo": previa.arquivo_nome},
                                usuario.id, usuario.perfil)

        con.execute(
            "INSERT INTO importacao (id, lote_id, arquivo_nome, hash_arquivo, importado_por,"
            " importado_em, novos, ignorados) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (novo_id(), d["lote_id"], previa.arquivo_nome, previa.hash_arquivo, usuario.id,
             momento, len(previa.novos), len(previa.itens) - len(previa.novos)))

    return ResultadoImportacao(len(previa.novos), len(previa.itens) - len(previa.novos),
                               previa.total_novos_centavos)


def historico_importacoes(con: sqlite3.Connection) -> list[sqlite3.Row]:
    return con.execute(
        "SELECT i.importado_em, i.arquivo_nome, i.novos, i.ignorados, u.nome AS supervisor"
        " FROM importacao i LEFT JOIN lote l ON l.id = i.lote_id"
        " LEFT JOIN usuario u ON u.id = l.gerado_por ORDER BY i.importado_em DESC LIMIT 50"
    ).fetchall()


def nome_sugerido(usuario: Usuario) -> str:
    base = "".join(ch if ch.isalnum() else "_" for ch in usuario.nome.split()[0].lower())
    return f"recibos_{base}_{date.today().isoformat()}.xlsx"
