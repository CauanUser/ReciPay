import base64
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from android_poc.recipay_xlsx import ler_previa


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def criar_xlsx(caminho: Path) -> None:
    comprovante = b"comprovante de teste sem dados pessoais"
    comp_hash = sha(comprovante)
    dados = {
        "versao_formato": 1,
        "tipo": "RECIBOS",
        "lote_id": "lote-teste-001",
        "gerado_por": {"id": "sup-1", "nome": "Supervisora Teste"},
        "gerado_em": "2026-10-26T10:00:00",
        "total_centavos": 12550,
        "colaboradores": [{"id": "col-1", "nome": "João Teste"}],
        "recibos": [{
            "id": "rec-1", "ref": "A1B2C3D4", "colaborador_id": "col-1",
            "valor_centavos": 12550, "comprovante": {"hash": comp_hash, "extensao": ".bin"}
        }],
        "cancelamentos": [],
    }
    texto = json.dumps(dados, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Recibos"
    ws.append(["Ref", "Data", "Colaborador", "Equipe", "Descrição", "Valor (R$)", "Categoria"])
    ws.append(["A1B2C3D4", "2026-10-26", "João Teste", "A", "Teste", 125.50, "Teste"])
    hidden = wb.create_sheet("_recipay")
    hidden.sheet_state = "veryHidden"
    hidden.append(["RECIPAY", 1, "Não altere"])
    hidden.append(["hash", 0, sha(texto.encode("utf-8"))])
    hidden.append(["dados", 0, texto])
    hidden.append([f"comp:{comp_hash}", 0, base64.b64encode(comprovante).decode("ascii")])
    wb.save(caminho)


class TestAndroidXlsxPreview(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "lote.xlsx"
        criar_xlsx(self.path)

    def test_valid_file_produces_preview_without_database_writes(self):
        preview = ler_previa(self.path)
        self.assertEqual(preview.supervisor, "Supervisora Teste")
        self.assertEqual(preview.lote_id, "lote-teste-001")
        self.assertEqual(preview.quantidade, 1)
        self.assertEqual(preview.total_centavos, 12550)
        self.assertEqual(preview.itens[0].colaborador, "João Teste")

    def test_modified_visible_amount_is_rejected(self):
        wb = load_workbook(self.path)
        wb["Recibos"]["F2"] = 999.99
        wb.save(self.path)
        with self.assertRaisesRegex(ValueError, "alterado"):
            ler_previa(self.path)

    def test_tampered_internal_json_is_rejected(self):
        wb = load_workbook(self.path)
        ws = wb["_recipay"]
        for row in ws.iter_rows():
            if row[0].value == "dados":
                row[2] = row[2].value.replace("12550", "22550")
        wb.save(self.path)
        with self.assertRaisesRegex(ValueError, "integridade"):
            ler_previa(self.path)

    def test_missing_technical_sheet_is_rejected(self):
        wb = load_workbook(self.path)
        del wb["_recipay"]
        wb.save(self.path)
        with self.assertRaisesRegex(ValueError, "abas"):
            ler_previa(self.path)

    def test_missing_or_damaged_attachment_is_rejected(self):
        wb = load_workbook(self.path)
        ws = wb["_recipay"]
        for row in ws.iter_rows():
            if str(row[0].value).startswith("comp:"):
                row[2] = base64.b64encode(b"alterado").decode("ascii")
        wb.save(self.path)
        with self.assertRaisesRegex(ValueError, "Comprovante"):
            ler_previa(self.path)


if __name__ == "__main__":
    unittest.main()
