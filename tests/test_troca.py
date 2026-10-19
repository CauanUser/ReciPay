"""Testes da troca de planilhas: supervisor exporta, gestor importa."""

import os
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from recipay import colaboradores, db, recibos, troca
from recipay.auth import criar_usuario_local


class TestTroca(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: __import__("shutil").rmtree(self.pasta, ignore_errors=True))
        self._usar_pc("supervisor")
        self.con_sup = db.conectar(":memory:")
        self.addCleanup(self.con_sup.close)
        self.sup = criar_usuario_local(self.con_sup, "Ana Supervisora", "SUPERVISOR", "123456")
        self.colab = colaboradores.criar(self.con_sup, self.sup, {
            "nome": "João Silva", "equipe": "Norte", "tipo_chave": "EMAIL",
            "chave_pix": "joao@x.com", "titular_pix": "João Silva"})
        foto = self.pasta / "nota.png"
        Image.new("RGB", (900, 1200), "white").save(foto)
        pdf = self.pasta / "nota.pdf"
        pdf.write_bytes(b"%PDF-1.4\n% teste\n%%EOF\n")
        hoje = date.today().isoformat()
        self.r1 = recibos.criar(self.con_sup, self.sup, {
            "colaborador_id": self.colab.id, "valor_centavos": 5590,
            "descricao": "Almoço", "data_despesa": hoje}, str(foto))
        self.r2 = recibos.criar(self.con_sup, self.sup, {
            "colaborador_id": self.colab.id, "valor_centavos": 12000,
            "descricao": "Combustível", "data_despesa": hoje}, str(pdf))
        self.arquivo = self.pasta / "recibos.xlsx"

    def _usar_pc(self, nome: str) -> None:
        """Cada 'PC' tem a sua pasta de dados (comprovantes)."""
        os.environ["RECIPAY_DADOS"] = str(self.pasta / nome)
        self.addCleanup(os.environ.pop, "RECIPAY_DADOS", None)

    def _exportar(self, ids=None):
        self._usar_pc("supervisor")
        return troca.exportar_recibos(self.con_sup, self.sup, ids or [self.r1.id, self.r2.id], self.arquivo)

    def _gestor(self):
        self._usar_pc("gestor")
        con = db.conectar(":memory:")
        self.addCleanup(con.close)
        return con, criar_usuario_local(con, "Bruno Gestor", "GESTOR", "123456")

    def test_exportar_marca_enviado_e_gera_abas(self):
        res = self._exportar()
        self.assertEqual((res.quantidade, res.total_centavos), (2, 17590))
        self.assertEqual({r.status for r in recibos.listar(self.con_sup)}, {"ENVIADO"})
        wb = load_workbook(self.arquivo)
        self.assertEqual(wb.sheetnames, ["Recibos", "Colaboradores", "Comprovantes", "_recipay"])
        self.assertEqual(wb["_recipay"].sheet_state, "veryHidden")
        with self.assertRaises(ValueError):  # exportado não edita mais
            recibos.atualizar(self.con_sup, self.sup, self.r1.id, {
                "colaborador_id": self.colab.id, "valor_centavos": 1, "descricao": "x x",
                "data_despesa": date.today().isoformat()})

    def test_importar_no_gestor(self):
        self._exportar()
        con, gestor = self._gestor()
        previa = troca.ler_arquivo(con, self.arquivo)
        self.assertEqual(previa.supervisor_nome, "Ana Supervisora")
        self.assertEqual(len(previa.novos), 2)
        self.assertEqual(len(previa.colaboradores_novos), 1)
        res = troca.importar(con, gestor, previa)
        self.assertEqual((res.novos, res.ignorados, res.total_centavos), (2, 0, 17590))

        lista = recibos.listar(con)
        self.assertEqual({r.status for r in lista}, {"PENDENTE"})
        self.assertEqual({r.supervisor_nome for r in lista}, {"Ana Supervisora"})
        self.assertEqual({r.saldo_centavos for r in lista}, {5590, 12000})
        # Comprovantes chegaram íntegros no "PC" do gestor
        from recipay import comprovantes
        for r in lista:
            self.assertTrue(comprovantes.caminho_absoluto(r.comprovante_arquivo).is_file())
        # Chave Pix disponível para o gestor pagar
        self.assertEqual(colaboradores.buscar(con, self.colab.id).chave_pix, "joao@x.com")

    def test_reimportar_com_outro_nome_nao_duplica(self):
        self._exportar()
        con, gestor = self._gestor()
        troca.importar(con, gestor, troca.ler_arquivo(con, self.arquivo))
        copia = self.pasta / "outro_nome.xlsx"
        copia.write_bytes(self.arquivo.read_bytes())
        previa = troca.ler_arquivo(con, copia)
        self.assertEqual(previa.novos, [])
        self.assertIsNotNone(previa.lote_ja_importado_em)
        res = troca.importar(con, gestor, previa)
        self.assertEqual((res.novos, res.ignorados), (0, 2))
        self.assertEqual(len(recibos.listar(con)), 2)

    def test_reenvio_so_traz_o_que_e_novo(self):
        self._exportar([self.r1.id])
        con, gestor = self._gestor()
        troca.importar(con, gestor, troca.ler_arquivo(con, self.arquivo))
        self._exportar([self.r1.id, self.r2.id])  # reenvia r1 junto com r2
        self._usar_pc("gestor")
        previa = troca.ler_arquivo(con, self.arquivo)
        self.assertEqual([i.dados["id"] for i in previa.novos], [self.r2.id])

    def test_planilha_alterada_no_excel_e_bloqueada(self):
        self._exportar()
        wb = load_workbook(self.arquivo)
        wb["Recibos"]["F2"] = 999.99
        wb.save(self.arquivo)
        con, _ = self._gestor()
        with self.assertRaisesRegex(ValueError, "alterada"):
            troca.ler_arquivo(con, self.arquivo)

    def test_dados_ocultos_alterados_sao_bloqueados(self):
        self._exportar()
        wb = load_workbook(self.arquivo)
        ws = wb["_recipay"]
        for linha in ws.iter_rows():
            if linha[0].value == "dados":
                linha[2].value = linha[2].value.replace("5590", "9590")
        wb.save(self.arquivo)
        con, _ = self._gestor()
        with self.assertRaisesRegex(ValueError, "alterada"):
            troca.ler_arquivo(con, self.arquivo)

    def test_planilha_qualquer_e_recusada(self):
        from openpyxl import Workbook
        outra = self.pasta / "qualquer.xlsx"
        Workbook().save(outra)
        con, _ = self._gestor()
        with self.assertRaisesRegex(ValueError, "não foi gerada pelo Recipay"):
            troca.ler_arquivo(con, outra)

    def test_chave_pix_alterada(self):
        self._exportar([self.r1.id])
        con, gestor = self._gestor()
        troca.importar(con, gestor, troca.ler_arquivo(con, self.arquivo))
        # Supervisor muda a chave e envia outro recibo
        self._usar_pc("supervisor")
        colaboradores.atualizar(self.con_sup, self.sup, self.colab.id, {
            "nome": "João Silva", "equipe": "Norte", "tipo_chave": "EMAIL",
            "chave_pix": "joao.novo@x.com", "titular_pix": "João Silva"})
        self._exportar([self.r2.id])
        self._usar_pc("gestor")
        previa = troca.ler_arquivo(con, self.arquivo)
        self.assertEqual(len(previa.chaves_alteradas), 1)
        troca.importar(con, gestor, previa, usar_chave_nova={self.colab.id: False})
        self.assertEqual(colaboradores.buscar(con, self.colab.id).chave_pix, "joao@x.com")

    def test_permissoes(self):
        self._exportar()
        con, gestor = self._gestor()
        with self.assertRaises(PermissionError):
            troca.exportar_recibos(con, gestor, [self.r1.id], self.arquivo)
        previa = troca.ler_arquivo(con, self.arquivo)
        with self.assertRaises(PermissionError):
            troca.importar(self.con_sup, self.sup, previa)


if __name__ == "__main__":
    unittest.main()
