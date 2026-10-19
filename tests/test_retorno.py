"""Testes do ciclo completo: supervisor → gestor → pagamento → retorno → supervisor."""

import os
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from recipay import colaboradores, db, pagamentos, recibos, retorno, troca
from recipay.auth import criar_usuario_local


class TestRetorno(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.pasta, True)
        self.addCleanup(os.environ.pop, "RECIPAY_DADOS", None)
        hoje = date.today().isoformat()

        # PC da supervisora Ana
        self._pc("ana")
        self.con_ana = db.conectar(":memory:")
        self.addCleanup(self.con_ana.close)
        self.ana = criar_usuario_local(self.con_ana, "Ana", "SUPERVISOR", "123456")
        foto = self.pasta / "nota.png"
        Image.new("RGB", (60, 60), "white").save(foto)
        self.joao = colaboradores.criar(self.con_ana, self.ana, {
            "nome": "João Silva", "equipe": "Norte", "tipo_chave": "EMAIL",
            "chave_pix": "joao@x.com", "titular_pix": "João Silva"})
        self.r1 = recibos.criar(self.con_ana, self.ana, {"colaborador_id": self.joao.id,
                                "valor_centavos": 10000, "descricao": "Almoço", "data_despesa": hoje}, str(foto))
        self.r2 = recibos.criar(self.con_ana, self.ana, {"colaborador_id": self.joao.id,
                                "valor_centavos": 5000, "descricao": "Táxi", "data_despesa": hoje}, str(foto))
        self.arq_recibos = self.pasta / "recibos.xlsx"
        troca.exportar_recibos(self.con_ana, self.ana, [self.r1.id, self.r2.id], self.arq_recibos)

        # PC do gestor importa e paga R$ 120 (r1 inteiro + 20 de r2)
        self._pc("gestor")
        self.con_g = db.conectar(":memory:")
        self.addCleanup(self.con_g.close)
        self.gestor = criar_usuario_local(self.con_g, "Bruno", "GESTOR", "123456")
        troca.importar(self.con_g, self.gestor, troca.ler_arquivo(self.con_g, self.arq_recibos))
        comp = self.pasta / "pix.pdf"
        comp.write_bytes(b"%PDF-1.4\npix 1\n%%EOF\n")
        self.tid = pagamentos.registrar(self.con_g, self.gestor, self.joao.id,
                                        {self.r1.id: 10000, self.r2.id: 2000}, hoje, str(comp))
        self.arq_ret = self.pasta / "retorno.xlsx"

    def _pc(self, nome):
        os.environ["RECIPAY_DADOS"] = str(self.pasta / nome)

    def _exportar_retorno(self, **kw):
        self._pc("gestor")
        return retorno.exportar_retorno(self.con_g, self.gestor, self.ana.id, self.arq_ret, **kw)

    def _situacoes_ana(self):
        return {r.id: (r.status, r.saldo_centavos) for r in recibos.listar(self.con_ana)}

    def test_ciclo_completo(self):
        sups = retorno.supervisores_para_retorno(self.con_g)
        self.assertEqual([(s.nome, s.pendentes, s.total_centavos) for s in sups], [("Ana", 1, 12000)])
        res = self._exportar_retorno()
        self.assertEqual((res.quantidade, res.total_centavos), (1, 12000))
        self.assertEqual(retorno.supervisores_para_retorno(self.con_g)[0].pendentes, 0)
        wb = load_workbook(self.arq_ret)
        self.assertEqual(wb.sheetnames, ["Pagamentos", "Comprovantes", "_recipay"])

        self._pc("ana")
        previa = retorno.ler_retorno(self.con_ana, self.ana, self.arq_ret)
        self.assertEqual((previa.gestor_nome, len(previa.novos)), ("Bruno", 1))
        r = retorno.importar_retorno(self.con_ana, self.ana, previa)
        self.assertEqual((r.novos, r.recibos_pagos, r.recibos_parciais), (1, 1, 1))
        self.assertEqual(self._situacoes_ana(), {self.r1.id: ("PAGO", 0), self.r2.id: ("PARCIAL", 3000)})
        t = pagamentos.listar_transferencias(self.con_ana)
        self.assertEqual([(x.valor_centavos, x.colaborador_nome) for x in t], [(12000, "João Silva")])

    def test_reimportar_nao_conta_duas_vezes(self):
        self._exportar_retorno()
        self._pc("ana")
        retorno.importar_retorno(self.con_ana, self.ana, retorno.ler_retorno(self.con_ana, self.ana, self.arq_ret))
        copia = self.pasta / "retorno_copia.xlsx"
        shutil.copy(self.arq_ret, copia)
        previa = retorno.ler_retorno(self.con_ana, self.ana, copia)
        self.assertEqual(previa.novos, [])
        retorno.importar_retorno(self.con_ana, self.ana, previa)
        self.assertEqual(self._situacoes_ana()[self.r2.id], ("PARCIAL", 3000))

    def test_sem_pagamentos_novos(self):
        self._exportar_retorno()
        with self.assertRaisesRegex(ValueError, "Não há pagamentos novos"):
            self._exportar_retorno()
        self._exportar_retorno(incluir_enviados=True)  # reenvio, se o arquivo se perdeu

    def test_estorno_bloqueado_depois_do_retorno(self):
        self._exportar_retorno()
        with self.assertRaisesRegex(ValueError, "já foi enviado"):
            pagamentos.estornar(self.con_g, self.gestor, self.tid, "teste")

    def test_retorno_de_outro_supervisor(self):
        self._exportar_retorno()
        self._pc("carla")
        con = db.conectar(":memory:")
        self.addCleanup(con.close)
        carla = criar_usuario_local(con, "Carla", "SUPERVISOR", "123456")
        with self.assertRaisesRegex(ValueError, "é para Ana"):
            retorno.ler_retorno(con, carla, self.arq_ret)

    def test_planilha_alterada(self):
        self._exportar_retorno()
        wb = load_workbook(self.arq_ret)
        wb["Pagamentos"]["F2"] = 1.0
        wb.save(self.arq_ret)
        self._pc("ana")
        with self.assertRaisesRegex(ValueError, "alterada"):
            retorno.ler_retorno(self.con_ana, self.ana, self.arq_ret)

    def test_planilha_de_recibos_nao_e_retorno(self):
        self._pc("ana")
        with self.assertRaisesRegex(ValueError, "não é uma planilha de retorno"):
            retorno.ler_retorno(self.con_ana, self.ana, self.arq_recibos)

    def test_permissoes(self):
        with self.assertRaises(PermissionError):
            retorno.exportar_retorno(self.con_ana, self.ana, self.ana.id, self.arq_ret)
        self._exportar_retorno()
        previa = retorno.ler_retorno(self.con_ana, self.ana, self.arq_ret)
        with self.assertRaises(PermissionError):
            retorno.importar_retorno(self.con_g, self.gestor, previa)


if __name__ == "__main__":
    unittest.main()
