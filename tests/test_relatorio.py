"""Categorias, cupom fiscal repetido, relatório mensal e séries dos gráficos."""

import os
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from recipay import colaboradores, db, pagamentos, recibos, relatorio, troca
from recipay.auth import criar_usuario_local


class TestCategoriasERelatorio(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.pasta, True)
        os.environ["RECIPAY_DADOS"] = str(self.pasta / "dados")
        self.addCleanup(os.environ.pop, "RECIPAY_DADOS", None)
        self.con = db.conectar(":memory:")
        self.addCleanup(self.con.close)
        self.sup = criar_usuario_local(self.con, "Ana", "SUPERVISOR", "123456", matricula="10")
        self.foto = self.pasta / "n.png"
        Image.new("RGB", (30, 30), "white").save(self.foto)
        self.joao = self._colab("João", "Norte", "j@x.com")
        self.maria = self._colab("Maria", "Sul", "m@x.com")
        self.hoje = date.today()

    def _colab(self, nome, equipe, chave):
        return colaboradores.criar(self.con, self.sup, {"nome": nome + " Silva", "equipe": equipe,
                                   "tipo_chave": "EMAIL", "chave_pix": chave, "titular_pix": nome})

    def _recibo(self, colab, valor, categoria, chave_nfe=None, data=None):
        return recibos.criar(self.con, self.sup, {
            "colaborador_id": colab.id, "valor_centavos": valor, "descricao": "Despesa", "categoria": categoria,
            "data_despesa": (data or self.hoje).isoformat(), "chave_nfe": chave_nfe}, str(self.foto))

    def test_categoria_padrao_e_invalida(self):
        r = recibos.criar(self.con, self.sup, {"colaborador_id": self.joao.id, "valor_centavos": 100,
                          "descricao": "Café", "data_despesa": self.hoje.isoformat()}, str(self.foto))
        self.assertEqual((r.categoria, r.categoria_nome), ("OUTROS", "Outros"))
        with self.assertRaisesRegex(ValueError, "categoria"):
            self._recibo(self.joao, 100, "XYZ")

    def test_mesmo_cupom_fiscal_duas_vezes(self):
        self._recibo(self.joao, 5000, "ALIMENTACAO", chave_nfe="1" * 44)
        with self.assertRaisesRegex(ValueError, "já foi usado"):
            self._recibo(self.maria, 5000, "ALIMENTACAO", chave_nfe="1" * 44)
        with self.assertRaisesRegex(ValueError, "Linha 2: repete o cupom"):
            recibos.criar_lote(self.con, self.sup, [
                ({"colaborador_id": self.joao.id, "valor_centavos": 1, "descricao": "aaa", "chave_nfe": "2" * 44,
                  "data_despesa": self.hoje.isoformat()}, str(self.foto)),
                ({"colaborador_id": self.joao.id, "valor_centavos": 1, "descricao": "aaa", "chave_nfe": "2" * 44,
                  "data_despesa": self.hoje.isoformat()}, str(self.foto))])

    def test_categoria_e_cupom_vao_para_o_gestor(self):
        r = self._recibo(self.joao, 5000, "COMBUSTIVEL", chave_nfe="3" * 44)
        arq = self.pasta / "a.xlsx"
        troca.exportar_recibos(self.con, self.sup, [r.id], arq)
        os.environ["RECIPAY_DADOS"] = str(self.pasta / "gestor")
        cg = db.conectar(":memory:")
        self.addCleanup(cg.close)
        g = criar_usuario_local(cg, "Bruno", "GESTOR", "123456")
        troca.importar(cg, g, troca.ler_arquivo(cg, arq))
        rg = recibos.buscar(cg, r.id)
        self.assertEqual((rg.categoria, rg.chave_nfe), ("COMBUSTIVEL", "3" * 44))

    def test_relatorio_do_mes(self):
        r1 = self._recibo(self.joao, 10000, "ALIMENTACAO")
        self._recibo(self.joao, 4000, "COMBUSTIVEL")
        self._recibo(self.maria, 6000, "ALIMENTACAO")
        cancelado = self._recibo(self.maria, 99900, "HOSPEDAGEM")
        recibos.cancelar(self.con, self.sup, cancelado.id, "Teste")
        # Simula o retorno de um pagamento parcial de r1
        self.con.execute("UPDATE recibo SET status = 'PENDENTE' WHERE status != 'CANCELADO'")
        self.con.execute("UPDATE usuario SET local = 0")
        g = criar_usuario_local(self.con, "Bruno", "GESTOR", "123456")
        comp = self.pasta / "pix.pdf"
        comp.write_bytes(b"%PDF-1.4\nx\n%%EOF\n")
        pagamentos.registrar(self.con, g, self.joao.id, {r1.id: 7000}, self.hoje.isoformat(), str(comp))

        rel = relatorio.do_mes(self.con, self.hoje.year, self.hoje.month)
        self.assertEqual((rel.geral.lancado, rel.geral.pago, rel.geral.pendente, rel.geral.quantidade),
                         (20000, 7000, 13000, 3))
        self.assertEqual(rel.por_categoria["ALIMENTACAO"].lancado, 16000)
        self.assertEqual(rel.por_equipe["Sul"].lancado, 6000)

        arq = self.pasta / "rel.xlsx"
        relatorio.exportar_excel(self.con, g, self.hoje.year, self.hoje.month, arq)
        wb = load_workbook(arq)
        self.assertEqual(wb.sheetnames, ["Resumo", "Por colaborador", "Recibos"])
        self.assertEqual(wb["Resumo"]["C5"].value, 200.0)   # lançado
        self.assertEqual(wb["Recibos"].max_row, 4)          # cabeçalho + 3 (sem o cancelado)
        self.assertEqual(len(wb["Resumo"]._charts), 2)

    def test_series_dos_graficos(self):
        self._recibo(self.joao, 1000, "OUTROS")
        self._recibo(self.maria, 3000, "OUTROS")
        meses = relatorio.ultimos_meses(self.con, 6)
        self.assertEqual(len(meses), 6)
        self.assertEqual(meses[-1][1], 4000)
        self.assertEqual(relatorio.top_equipes(self.con), [("Sul", 3000), ("Norte", 1000)])
        self.assertIn((self.hoje.year, self.hoje.month), relatorio.meses_com_dados(self.con))


if __name__ == "__main__":
    unittest.main()
