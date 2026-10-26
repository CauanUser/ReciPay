"""Leitor independente de prévia XLSX para o protótipo Android.

Este módulo não grava no banco nem substitui recipay.troca. Ele valida a
estrutura técnica e retorna apenas um resumo para a tela de prévia.
"""
from __future__ import annotations

import base64
import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

MARCA = "RECIPAY"
ABA_DADOS = "_recipay"
VERSAO_FORMATO = 1


@dataclass(frozen=True)
class ItemResumo:
    referencia: str
    colaborador: str
    valor_centavos: int
    comprovante_ok: bool


@dataclass(frozen=True)
class PreviaXlsx:
    arquivo: str
    versao_formato: int
    lote_id: str
    supervisor: str
    itens: tuple[ItemResumo, ...]
    total_centavos: int

    @property
    def quantidade(self) -> int:
        return len(self.itens)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def ler_previa(caminho: str | Path) -> PreviaXlsx:
    """Valida arquivo de recibos e retorna resumo sem gravar dados.

    Rejeita arquivos externos ao Recipay, versão futura, hash inconsistente,
    comprovantes ausentes/danificados e valores visíveis alterados.
    """
    caminho = Path(caminho)
    try:
        wb = load_workbook(caminho, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Não foi possível abrir a planilha: {exc}") from exc

    try:
        if ABA_DADOS not in wb.sheetnames or "Recibos" not in wb.sheetnames:
            raise ValueError("Arquivo não reconhecido: abas do Recipay ausentes.")
        linhas = wb[ABA_DADOS].iter_rows(values_only=True)
        cab = next(linhas, None)
        if not cab or len(cab) < 2 or cab[0] != MARCA:
            raise ValueError("Planilha do Recipay danificada (marca não encontrada).")
        versao = int(cab[1] or 0)
        if versao > VERSAO_FORMATO:
            raise ValueError("Arquivo usa versão mais nova do formato Recipay.")
        partes: dict[str, list[tuple[int, str]]] = {}
        for linha in linhas:
            if not linha or len(linha) < 3 or linha[0] is None:
                continue
            try:
                indice = int(linha[1])
            except (TypeError, ValueError) as exc:
                raise ValueError("Fragmento técnico inválido.") from exc
            partes.setdefault(str(linha[0]), []).append((indice, str(linha[2] or "")))
        unidos = {
            chave: "".join(texto for _, texto in sorted(fragmentos))
            for chave, fragmentos in partes.items()
        }
    finally:
        wb.close()

    texto = unidos.get("dados")
    if not texto:
        raise ValueError("Planilha danificada: dados técnicos ausentes.")
    if _sha256(texto.encode("utf-8")) != unidos.get("hash"):
        raise ValueError("A integridade dos dados não confere; solicite novo arquivo.")
    try:
        dados: dict[str, Any] = json.loads(texto)
    except (TypeError, json.JSONDecodeError) as exc:
        raise ValueError("JSON interno inválido.") from exc
    if dados.get("tipo") != "RECIBOS":
        raise ValueError("Este protótipo aceita apenas arquivos de recibos.")
    if int(dados.get("versao_formato", 0)) > VERSAO_FORMATO:
        raise ValueError("Arquivo usa versão mais nova do formato Recipay.")

    comprovantes: dict[str, bytes] = {}
    for chave, texto_b64 in unidos.items():
        if chave.startswith("comp:"):
            try:
                comprovantes[chave[5:]] = base64.b64decode(texto_b64, validate=True)
            except (ValueError, TypeError) as exc:
                raise ValueError("Comprovante codificado de forma inválida.") from exc

    ws = load_workbook(caminho, read_only=True, data_only=True)
    try:
        visiveis: dict[str, set[int]] = {}
        for linha in ws["Recibos"].iter_rows(min_row=2, values_only=True):
            if linha and len(linha) > 5 and linha[0] and isinstance(linha[5], (int, float)):
                visiveis.setdefault(str(linha[0]), set()).add(round(linha[5] * 100))
    finally:
        ws.close()

    colabs = {str(c["id"]): c for c in dados.get("colaboradores", [])}
    itens: list[ItemResumo] = []
    for item in dados.get("recibos", []):
        try:
            valor = int(item["valor_centavos"])
            ref = str(item["ref"])
            comp_hash = str(item["comprovante"]["hash"])
            colab = colabs[str(item["colaborador_id"])]
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError("Dados de recibo incompletos ou inconsistentes.") from exc
        if valor <= 0:
            raise ValueError(f"Valor inválido no recibo {ref}.")
        if visiveis.get(ref) != {valor}:
            raise ValueError(f"Valor visível do recibo {ref} foi alterado ou está ausente.")
        conteudo = comprovantes.get(comp_hash)
        if conteudo is None or _sha256(conteudo) != comp_hash:
            raise ValueError(f"Comprovante do recibo {ref} ausente ou danificado.")
        itens.append(ItemResumo(ref, str(colab.get("nome", "")), valor, True))

    total = sum(i.valor_centavos for i in itens)
    if total != int(dados.get("total_centavos", -1)):
        raise ValueError("O total dos recibos não confere com o total do lote.")
    supervisor = dados.get("gerado_por", {}).get("nome")
    lote_id = dados.get("lote_id")
    if not supervisor or not lote_id:
        raise ValueError("Identificação do supervisor ou lote ausente.")
    return PreviaXlsx(
        arquivo=caminho.name,
        versao_formato=versao,
        lote_id=str(lote_id),
        supervisor=str(supervisor),
        itens=tuple(itens),
        total_centavos=total,
    )
